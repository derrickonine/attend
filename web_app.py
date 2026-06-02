import csv
import io
import json
import threading
import time
import webbrowser
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from flask import Flask, Response, jsonify, request, send_file, send_from_directory
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
except ImportError:
    A4 = None
    colors = None
    cm = None
    SimpleDocTemplate = None
    Table = None
    TableStyle = None
    Paragraph = None
    Spacer = None
    getSampleStyleSheet = None
    ParagraphStyle = None

from app import (
    APP_DIR,
    EXPORTS_DIR,
    FACES_DIR,
    Database,
    crop_face,
    descriptor_distance,
    detector_status,
    detect_faces,
    face_descriptor,
)


WEB_DIR = APP_DIR / "web"
app = Flask(__name__, static_folder=str(WEB_DIR), static_url_path="")
db = Database()
db.close_open_sessions("Fermee automatiquement au demarrage de l'application web")


class CameraSession:
    def __init__(self):
        self.lock = threading.Lock()
        self.cap = None
        self.mode = "idle"
        self.camera_index = 0
        self.student = None
        self.course = None
        self.session_id = None
        self.target_count = 8
        self.captured = 0
        self.stable_frames = 0
        self.last_capture_at = 0.0
        self.known_encodings = []
        self.known_metadata = []
        self.marked = set()
        self.logs = []
        self.last_status = "Camera inactive"
        self.last_detected = 0
        self.last_fps = 0
        self.last_distance = None
        self.recognized_count = 0
        self.confidence_scores = []
        self.last_recognized = None
        # Live presence list: [{student_id, name, time, confidence}]
        self.live_presence = []

    def add_log(self, message, kind="info"):
        with self.lock:
            self.logs.insert(
                0,
                {
                    "time": datetime.now().strftime("%H:%M:%S"),
                    "message": message,
                    "kind": kind,
                },
            )
            self.logs = self.logs[:120]

    def open(self, camera_index):
        if self.cap is not None:
            self.cap.release()
        self.cap = cv2.VideoCapture(int(camera_index))
        if not self.cap.isOpened():
            self.cap = None
            raise RuntimeError("Impossible d'ouvrir la camera")
        self.camera_index = int(camera_index)

    def stop(self):
        with self.lock:
            if self.cap is not None:
                self.cap.release()
            if self.mode == "attendance" and self.session_id:
                db.close_session(self.session_id)
            self.cap = None
            self.mode = "idle"
            self.student = None
            self.course = None
            self.session_id = None
            self.known_encodings = []
            self.known_metadata = []
            self.marked = set()
            self.live_presence = []
            self.recognized_count = 0
            self.confidence_scores = []
            self.last_recognized = None
            self.last_status = "Camera inactive"

    def read(self):
        if self.cap is None:
            return None
        ok, frame = self.cap.read()
        return frame if ok else None


camera = CameraSession()


def row_to_dict(row):
    return dict(row) if row is not None else None


def get_teachers_payload():
    teachers = db.conn.execute("SELECT * FROM teachers ORDER BY full_name").fetchall()
    payload = []
    for teacher in teachers:
        courses = [row_to_dict(course) for course in db.get_teacher_courses(teacher["id"])]
        payload.append({**row_to_dict(teacher), "courses": courses})
    return payload


def get_classes_payload():
    rows = db.conn.execute("SELECT * FROM classes ORDER BY name").fetchall()
    return [row_to_dict(row) for row in rows]


def get_students_payload(class_id=None):
    if class_id:
        rows = db.list_students(class_id)
    else:
        rows = db.conn.execute(
            """
            SELECT students.*,
                   classes.name AS class_name,
                   COUNT(face_samples.id) AS samples_count,
                   COUNT(face_samples.encoding) AS encodings_count,
                   MAX(face_samples.created_at) AS last_sample_at
            FROM students
            JOIN classes ON classes.id = students.class_id
            LEFT JOIN face_samples ON face_samples.student_id = students.id
            GROUP BY students.id
            ORDER BY classes.name, students.full_name
            """
        ).fetchall()
    payload = []
    for row in rows:
        item = row_to_dict(row)
        count = int(item.get("samples_count") or 0)
        if count >= 20:
            quality = "Excellent"
        elif count >= 15:
            quality = "Bon"
        else:
            quality = "Faible"
        item["dataset_quality"] = quality
        item["recognition_ready"] = count >= 15 and bool(item.get("biometric_consent"))
        payload.append(item)
    return payload


def get_sessions_payload(class_id=None):
    if class_id:
        rows = db.list_sessions(class_id)
    else:
        rows = db.conn.execute(
            """
            SELECT attendance_sessions.*,
                   classes.name AS class_name,
                   courses.name AS course_name,
                   teachers.full_name AS teacher_name,
                   COUNT(students.id) AS total_count,
                   SUM(CASE WHEN attendance_records.status = 'present' THEN 1 ELSE 0 END) AS present_count
            FROM attendance_sessions
            JOIN classes ON classes.id = attendance_sessions.class_id
            JOIN courses ON courses.id = attendance_sessions.course_id
            JOIN teachers ON teachers.id = attendance_sessions.teacher_id
            JOIN students ON students.class_id = attendance_sessions.class_id
            LEFT JOIN attendance_records ON attendance_records.session_id = attendance_sessions.id
             AND attendance_records.student_id = students.id
            GROUP BY attendance_sessions.id
            ORDER BY attendance_sessions.started_at DESC
            """
        ).fetchall()
    return [row_to_dict(row) for row in rows]


def get_system_stats_payload():
    sessions = get_sessions_payload()
    students = get_students_payload()
    total_sessions = len(sessions)
    total_students = len(students)
    rates = []
    for session in sessions:
        total = int(session.get("total_count") or total_students or 0)
        present = int(session.get("present_count") or 0)
        rates.append(round((present / total) * 100, 1) if total else 0)
    avg_rate = round(sum(rates) / len(rates), 1) if rates else 0
    return {
        "total_sessions": total_sessions,
        "total_students": total_students,
        "total_classes": len(get_classes_payload()),
        "total_teachers": len(get_teachers_payload()),
        "average_attendance_rate": avg_rate,
        "best_attendance_rate": max(rates) if rates else 0,
        "sessions_trend": [
            {
                "id": session["id"],
                "label": f"#{session['id']}",
                "rate": rates[index] if index < len(rates) else 0,
                "present": int(session.get("present_count") or 0),
                "total": int(session.get("total_count") or total_students or 0),
            }
            for index, session in enumerate(sessions[:8])
        ],
    }


def get_course(course_id):
    row = db.conn.execute(
        """
        SELECT courses.*, classes.name AS class_name, teachers.full_name AS teacher_name
        FROM courses
        JOIN classes ON classes.id = courses.class_id
        JOIN teachers ON teachers.id = courses.teacher_id
        WHERE courses.id = ?
        """,
        (course_id,),
    ).fetchone()
    return row


def get_session(session_id):
    return db.conn.execute(
        """
        SELECT attendance_sessions.*, classes.name AS class_name,
               courses.name AS course_name, teachers.full_name AS teacher_name
        FROM attendance_sessions
        JOIN classes ON classes.id = attendance_sessions.class_id
        JOIN courses ON courses.id = attendance_sessions.course_id
        JOIN teachers ON teachers.id = attendance_sessions.teacher_id
        WHERE attendance_sessions.id = ?
        """,
        (session_id,),
    ).fetchone()


def build_export_payload(session, rows):
    presents = sum(1 for row in rows if row["status"] == "present")
    return {
        "session": {
            "id": session["id"],
            "class": session["class_name"],
            "course": session["course_name"],
            "teacher": session["teacher_name"],
            "started_at": session["started_at"],
            "ended_at": session["ended_at"],
        },
        "summary": {
            "students": len(rows),
            "present": presents,
            "absent": len(rows) - presents,
            "rate": round((presents / len(rows)) * 100, 1) if rows else 0,
        },
        "rows": [
            {
                "student_id": row["student_id"],
                "full_name": row["full_name"],
                "status": row["status"],
                "recognized_at": row["recognized_at"],
                "confidence": row["confidence"],
            }
            for row in rows
        ],
    }


def draw_idle_frame(width=960, height=540):
    frame = np.zeros((height, width, 3), dtype=np.uint8)
    frame[:] = (28, 24, 20)
    cv2.putText(frame, "FaceAttend - camera inactive", (40, height // 2), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (230, 230, 230), 2)
    return frame


def save_student_face(frame, face_box):
    student = camera.student
    class_name = student["class_name"]
    student_dir = FACES_DIR / class_name / safe_name(student["full_name"])
    student_dir.mkdir(parents=True, exist_ok=True)
    face_img = crop_face(frame, face_box)
    descriptor = face_descriptor(face_img)
    image_path = student_dir / f"face_{int(time.time())}_{camera.captured + 1}.jpg"
    if not cv2.imwrite(str(image_path), face_img):
        camera.last_status = "Erreur: image non enregistree"
        return
    db.add_face_sample(student["id"], image_path, descriptor)
    camera.captured += 1
    camera.last_capture_at = time.time()
    camera.stable_frames = 0
    camera.last_status = f"Photo {camera.captured}/{camera.target_count} enregistree"
    camera.add_log(camera.last_status, "success")


def handle_capture_frame(frame, faces):
    if not camera.student:
        return
    if len(faces) == 1:
        camera.stable_frames += 1
        camera.last_status = "Visage stable detecte"
        if camera.stable_frames >= 10 and time.time() - camera.last_capture_at >= 1.0:
            save_student_face(frame, faces[0])
    else:
        camera.stable_frames = 0
        camera.last_status = "Garde un seul visage visible"
    if camera.captured >= camera.target_count:
        camera.last_status = "Capture terminee"
        camera.add_log(f"Capture terminee pour {camera.student['full_name']}", "success")


def handle_attendance_frame(frame, faces):
    if not camera.known_encodings:
        camera.last_status = "Aucun visage etudiant enregistre"
        return
    threshold = 0.36
    for face_box in faces:
        face_img = crop_face(frame, face_box)
        descriptor = face_descriptor(face_img)
        distances = [descriptor_distance(known, descriptor) for known in camera.known_encodings]
        if not distances:
            continue
        best_index = int(np.argmin(distances))
        distance = float(distances[best_index])
        camera.last_distance = round(distance, 3)
        if distance <= threshold:
            student = camera.known_metadata[best_index]
            confidence = max(0.0, 1.0 - distance)
            if student["student_id"] not in camera.marked:
                db.mark_present(camera.session_id, student["student_id"], confidence)
                camera.marked.add(student["student_id"])
                now_str = datetime.now().strftime("%H:%M:%S")
                camera.recognized_count += 1
                camera.confidence_scores.append(confidence)
                camera.confidence_scores = camera.confidence_scores[-50:]
                camera.last_recognized = student["name"]
                # Add to live presence list
                camera.live_presence.append({
                    "student_id": student["student_id"],
                    "name": student["name"],
                    "time": now_str,
                    "confidence": round(confidence * 100)
                })
                camera.add_log(f"Presence: {student['name']} ({round(confidence * 100)}%)", "success")
            x, y, w, h = face_box
            cv2.putText(frame, student["name"], (x, max(26, y - 10)), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 190, 90), 2)
        else:
            x, y, w, h = face_box
            cv2.putText(frame, "Inconnu", (x, max(26, y - 10)), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (40, 120, 255), 2)
    camera.last_status = f"{len(camera.marked)} present(s)"


def annotate_frame(frame):
    t0 = time.time()
    faces = detect_faces(frame)
    camera.last_detected = len(faces)
    for face_box in faces:
        x, y, w, h = face_box
        color = (0, 190, 90) if camera.mode in ("capture", "attendance") else (200, 120, 40)
        cv2.rectangle(frame, (x, y), (x + w, y + h), color, 2)

    if camera.mode == "capture":
        handle_capture_frame(frame, faces)
        title = f"Inscription: {camera.student['full_name']} - {camera.captured}/{camera.target_count}"
    elif camera.mode == "attendance":
        handle_attendance_frame(frame, faces)
        title = f"Appel: {camera.course['class_name']} - presents {len(camera.marked)}"
    else:
        title = "Camera active"

    camera.last_fps = round(1 / max(time.time() - t0, 0.001), 1)
    cv2.rectangle(frame, (0, 0), (frame.shape[1], 82), (22, 20, 18), -1)
    cv2.putText(frame, title, (18, 32), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (245, 245, 245), 2)
    cv2.putText(frame, camera.last_status, (18, 64), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (210, 210, 210), 2)


def frame_generator():
    while True:
        frame = camera.read()
        if frame is None:
            frame = draw_idle_frame()
        else:
            annotate_frame(frame)
        ok, buffer = cv2.imencode(".jpg", frame, [int(cv2.IMWRITE_JPEG_QUALITY), 82])
        if not ok:
            continue
        yield b"--frame\r\nContent-Type: image/jpeg\r\n\r\n" + buffer.tobytes() + b"\r\n"
        time.sleep(0.03)


def safe_name(value):
    return "".join(ch if ch.isalnum() or ch in (" ", "-", "_") else "_" for ch in value).strip()


@app.get("/")
def index():
    return send_from_directory(WEB_DIR, "index.html")


@app.get("/api/bootstrap")
def bootstrap():
    class_id = request.args.get("class_id", type=int)
    return jsonify(
        {
            "teachers": get_teachers_payload(),
            "classes": get_classes_payload(),
            "students": get_students_payload(class_id),
            "sessions": get_sessions_payload(class_id),
            "camera": get_camera_status_payload(),
            "stats": get_system_stats_payload(),
        }
    )


@app.post("/api/students")
def add_student():
    data = request.get_json(force=True)
    full_name = data.get("full_name", "").strip()
    class_id = int(data.get("class_id") or 0)
    biometric_consent = bool(data.get("biometric_consent"))
    if not full_name or not class_id:
        return jsonify({"error": "Nom et classe requis"}), 400
    if not biometric_consent:
        return jsonify({"error": "Autorisation des donnees biometriques requise"}), 400
    student = db.get_student_or_create(full_name, class_id, biometric_consent)
    class_row = db.conn.execute("SELECT name FROM classes WHERE id = ?", (class_id,)).fetchone()
    camera.add_log(f"Etudiant ajoute: {full_name}", "success")
    return jsonify(
        {
            "student": {
                **row_to_dict(student),
                "class_name": class_row["name"],
                "samples_count": 0,
                "encodings_count": 0,
                "dataset_quality": "Faible",
                "recognition_ready": False,
            }
        }
    )


@app.get("/api/students/<int:student_id>/photos")
def student_photos(student_id):
    rows = db.student_face_samples(student_id)
    return jsonify(
        {
            "photos": [
                {
                    "id": row["id"],
                    "image_path": row["image_path"],
                    "created_at": row["created_at"],
                    "url": f"/api/face-samples/{row['id']}/image",
                }
                for row in rows
            ]
        }
    )


@app.get("/api/face-samples/<int:sample_id>/image")
def face_sample_image(sample_id):
    row = db.conn.execute("SELECT image_path FROM face_samples WHERE id = ?", (sample_id,)).fetchone()
    if not row:
        return jsonify({"error": "Image introuvable"}), 404
    path = APP_DIR / row["image_path"]
    if not path.exists():
        return jsonify({"error": "Fichier image introuvable"}), 404
    return send_file(path)


@app.post("/api/students/<int:student_id>/regenerate")
def regenerate_student(student_id):
    student = db.conn.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    if not student:
        return jsonify({"error": "Etudiant introuvable"}), 404
    updated = db.regenerate_student_encodings(student_id)
    camera.add_log(f"Encodages regeneres pour {student['full_name']}: {updated}", "success")
    return jsonify({"ok": True, "updated": updated, "students": get_students_payload()})


@app.delete("/api/students/<int:student_id>")
def delete_student(student_id):
    student = db.conn.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    if not student:
        return jsonify({"error": "Etudiant introuvable"}), 404
    db.delete_student(student_id)
    camera.add_log(f"Etudiant supprime: {student['full_name']}", "warn")
    return jsonify({"ok": True, "students": get_students_payload()})


@app.post("/api/classes")
def add_class():
    data = request.get_json(force=True)
    name = data.get("name", "").strip().upper()
    level = data.get("level", "").strip()
    school_year = data.get("school_year", "2025-2026").strip()
    if not name:
        return jsonify({"error": "Nom de classe requis"}), 400
    db.conn.execute(
        "INSERT OR IGNORE INTO classes(name, level, school_year) VALUES (?, ?, ?)",
        (name, level, school_year),
    )
    db.conn.commit()
    camera.add_log(f"Classe ajoutee: {name}", "success")
    return jsonify({"classes": get_classes_payload()})


@app.post("/api/teachers")
def add_teacher():
    data = request.get_json(force=True)
    full_name = data.get("full_name", "").strip()
    teacher_code = data.get("teacher_code", "").strip().upper()
    course_name = data.get("course_name", "").strip()
    class_id = int(data.get("class_id") or 0)
    weekday = data.get("weekday", "").strip()
    starts_at = data.get("starts_at", "").strip()
    ends_at = data.get("ends_at", "").strip()
    if not full_name or not teacher_code:
        return jsonify({"error": "Nom et ID professeur requis"}), 400
    db.conn.execute(
        "INSERT OR IGNORE INTO teachers(teacher_code, full_name) VALUES (?, ?)",
        (teacher_code, full_name),
    )
    teacher = db.get_teacher_by_code(teacher_code)
    if course_name and class_id:
        db.conn.execute(
            """
            INSERT OR IGNORE INTO courses(name, class_id, teacher_id, weekday, starts_at, ends_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (course_name, class_id, teacher["id"], weekday, starts_at, ends_at),
        )
    db.conn.commit()
    camera.add_log(f"Professeur ajoute: {full_name}", "success")
    return jsonify({"teachers": get_teachers_payload()})


@app.post("/api/courses")
def add_course():
    data = request.get_json(force=True)
    name = data.get("name", "").strip()
    class_id = int(data.get("class_id") or 0)
    teacher_id = int(data.get("teacher_id") or 0)
    weekday = data.get("weekday", "").strip()
    starts_at = data.get("starts_at", "").strip()
    ends_at = data.get("ends_at", "").strip()
    if not name or not class_id or not teacher_id:
        return jsonify({"error": "Nom, classe et professeur requis"}), 400
    db.conn.execute(
        """
        INSERT OR IGNORE INTO courses(name, class_id, teacher_id, weekday, starts_at, ends_at)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (name, class_id, teacher_id, weekday, starts_at, ends_at),
    )
    db.conn.commit()
    camera.add_log(f"Cours ajoute: {name}", "success")
    return jsonify({"teachers": get_teachers_payload()})


@app.post("/api/capture/start")
def start_capture():
    data = request.get_json(force=True)
    student_id = int(data.get("student_id") or 0)
    camera_index = int(data.get("camera_index") or 0)
    target_count = int(data.get("target_count") or 15)
    student = db.conn.execute(
        """
        SELECT students.*, classes.name AS class_name
        FROM students
        JOIN classes ON classes.id = students.class_id
        WHERE students.id = ?
        """,
        (student_id,),
    ).fetchone()
    if not student:
        return jsonify({"error": "Etudiant introuvable"}), 404
    if camera.mode != "idle":
        camera.stop()
    try:
        camera.open(camera_index)
    except RuntimeError as exc:
        return jsonify({"error": str(exc)}), 400
    with camera.lock:
        camera.mode = "capture"
        camera.student = row_to_dict(student)
        camera.target_count = target_count
        camera.captured = 0
        camera.stable_frames = 0
        camera.last_capture_at = 0.0
        camera.last_status = "Capture prete"
    camera.add_log(f"Capture visage lancee pour {student['full_name']}", "info")
    return jsonify({"ok": True})


@app.post("/api/attendance/start")
def start_attendance():
    data = request.get_json(force=True)
    course_id = int(data.get("course_id") or 0)
    camera_index = int(data.get("camera_index") or 0)
    course = get_course(course_id)
    if not course:
        return jsonify({"error": "Cours introuvable"}), 404
    known_encodings, known_metadata = db.load_known_faces(course["class_id"])
    if not known_encodings:
        return jsonify({"error": "Aucun visage inscrit pour cette classe"}), 400
    if camera.mode != "idle":
        camera.stop()
    db.close_open_sessions("Fermee automatiquement avant une nouvelle seance")
    try:
        camera.open(camera_index)
    except RuntimeError as exc:
        return jsonify({"error": str(exc)}), 400
    session_id = db.create_session(course)
    with camera.lock:
        camera.mode = "attendance"
        camera.course = row_to_dict(course)
        camera.session_id = session_id
        camera.known_encodings = known_encodings
        camera.known_metadata = known_metadata
        camera.marked = set()
        camera.live_presence = []
        camera.last_status = "Appel en cours"
    camera.add_log(f"Appel lance: {course['class_name']} - {course['name']}", "info")
    return jsonify({"ok": True, "session_id": session_id})


@app.post("/api/camera/stop")
def stop_camera():
    camera.stop()
    camera.add_log("Camera arretee", "warn")
    return jsonify({"ok": True})


def get_camera_status_payload():
    with camera.lock:
        return {
            "mode": camera.mode,
            "captured": camera.captured,
            "target_count": camera.target_count,
            "status": camera.last_status,
            "detected": camera.last_detected,
            "fps": camera.last_fps,
            "distance": camera.last_distance,
            "marked": len(camera.marked),
            "session_id": camera.session_id,
            "detector": detector_status(),
            "camera_index": camera.camera_index,
            "recognized": camera.recognized_count,
            "avg_confidence": round(sum(camera.confidence_scores) / len(camera.confidence_scores), 3)
            if camera.confidence_scores
            else None,
            "last_recognized": camera.last_recognized,
            "logs": list(camera.logs[:80]),
            "live_presence": list(camera.live_presence),
        }


@app.get("/api/camera/status")
def camera_status():
    return jsonify(get_camera_status_payload())


@app.get("/video_feed")
def video_feed():
    return Response(frame_generator(), mimetype="multipart/x-mixed-replace; boundary=frame")


@app.get("/api/sessions/<int:session_id>/report")
def session_report(session_id):
    return jsonify({"rows": [row_to_dict(row) for row in db.attendance_report_rows(session_id)]})


@app.post("/api/sessions/<int:session_id>/attendance/<int:student_id>")
def update_attendance_status(session_id, student_id):
    data = request.get_json(force=True)
    status = data.get("status", "").strip().lower()
    if status not in ("present", "absent"):
        return jsonify({"error": "Statut non supporte"}), 400
    session = get_session(session_id)
    if not session:
        return jsonify({"error": "Seance introuvable"}), 404
    student = db.conn.execute(
        "SELECT * FROM students WHERE id = ? AND class_id = ?",
        (student_id, session["class_id"]),
    ).fetchone()
    if not student:
        return jsonify({"error": "Etudiant introuvable dans cette classe"}), 404

    db.set_attendance_status(session_id, student_id, status, 1.0 if status == "present" else None)
    with camera.lock:
        if camera.session_id == session_id:
            if status == "present":
                camera.marked.add(student_id)
                if not any(item["student_id"] == student_id for item in camera.live_presence):
                    camera.live_presence.append(
                        {
                            "student_id": student_id,
                            "name": student["full_name"],
                            "time": datetime.now().strftime("%H:%M:%S"),
                            "confidence": 100,
                        }
                    )
            else:
                camera.marked.discard(student_id)
                camera.live_presence = [
                    item for item in camera.live_presence if item["student_id"] != student_id
                ]
    camera.add_log(f"Correction manuelle: {student['full_name']} -> {status}", "warn")
    return jsonify({"ok": True, "rows": [row_to_dict(row) for row in db.attendance_report_rows(session_id)]})


@app.get("/api/export/<int:session_id>/<fmt>")
def export_session(session_id, fmt):
    session = get_session(session_id)
    if not session:
        return jsonify({"error": "Seance introuvable"}), 404
    rows = db.attendance_report_rows(session_id)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"presence_{session['class_name']}_session_{session_id}_{stamp}.{fmt}"
    path = EXPORTS_DIR / filename
    presents = sum(1 for r in rows if r["status"] == "present")
    absents = len(rows) - presents
    rate = f"{round(presents / len(rows) * 100)}%" if rows else "0%"

    if fmt == "csv":
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["=== FEUILLE DE PRESENCE ==="])
            writer.writerow(["Classe", session["class_name"], "Cours", session["course_name"]])
            writer.writerow(["Professeur", session["teacher_name"], "Date session", session["started_at"]])
            writer.writerow(["Presents", presents, "Absents", absents, "Taux", rate])
            writer.writerow([])
            writer.writerow(["N°", "Etudiant", "Statut", "Heure de reconnaissance", "Confiance (%)"])
            for i, row in enumerate(rows, 1):
                conf = f"{round(float(row['confidence']) * 100)}%" if row['confidence'] else "-"
                status = "PRESENT" if row["status"] == "present" else "ABSENT"
                writer.writerow([i, row["full_name"], status, row["recognized_at"] or "-", conf])
            writer.writerow([])
            writer.writerow(["TOTAL", "", f"Presents: {presents}", f"Absents: {absents}", f"Taux: {rate}"])

    elif fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Feuille de Presence"
        
        # Styles
        header_fill = PatternFill("solid", fgColor="C65F2A")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        title_font = Font(bold=True, size=14, color="1B1713")
        present_fill = PatternFill("solid", fgColor="E7F5ED")
        absent_fill = PatternFill("solid", fgColor="FDEBE8")
        border = Border(
            left=Side(style='thin', color='E2DBD1'),
            right=Side(style='thin', color='E2DBD1'),
            top=Side(style='thin', color='E2DBD1'),
            bottom=Side(style='thin', color='E2DBD1')
        )
        center = Alignment(horizontal='center', vertical='center')
        
        # Title
        ws.merge_cells("A1:E1")
        ws["A1"] = "FEUILLE DE PRESENCE"
        ws["A1"].font = title_font
        ws["A1"].alignment = center
        ws.row_dimensions[1].height = 30
        
        # Info rows
        ws["A2"] = "Classe:"
        ws["B2"] = session["class_name"]
        ws["A3"] = "Cours:"
        ws["B3"] = session["course_name"]
        ws["A4"] = "Professeur:"
        ws["B4"] = session["teacher_name"]
        ws["A5"] = "Date:"
        ws["B5"] = session["started_at"]
        ws["A6"] = "Synthese:"
        ws["B6"] = f"Presents: {presents} | Absents: {absents} | Taux: {rate}"
        for r in range(2, 7):
            ws[f"A{r}"].font = Font(bold=True)
        
        ws.row_dimensions[7].height = 8
        
        # Table headers
        headers = ["N°", "Nom de l'etudiant", "Statut", "Heure reconnaissance", "Confiance"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=8, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border
        ws.row_dimensions[8].height = 22
        
        # Data rows
        for i, row in enumerate(rows, 1):
            r = i + 8
            conf = f"{round(float(row['confidence']) * 100)}%" if row['confidence'] else "-"
            status = "PRESENT" if row["status"] == "present" else "ABSENT"
            vals = [i, row["full_name"], status, row["recognized_at"] or "-", conf]
            fill = present_fill if row["status"] == "present" else absent_fill
            for col, val in enumerate(vals, 1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.fill = fill
                cell.border = border
                if col in (1, 3, 5):
                    cell.alignment = center
        
        # Summary
        last_r = len(rows) + 10
        ws.cell(row=last_r, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=last_r, column=2, value=f"Presents: {presents}  |  Absents: {absents}  |  Taux: {rate}")
        
        # Column widths
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 22
        ws.column_dimensions['E'].width = 14
        
        wb.save(path)

    elif fmt == "pdf":
        if SimpleDocTemplate is None:
            return jsonify({"error": "Export PDF indisponible: installe reportlab avec pip install reportlab"}), 400
        doc = SimpleDocTemplate(str(path), pagesize=A4,
                                rightMargin=1.5*cm, leftMargin=1.5*cm,
                                topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()
        story = []
        
        # Title
        title_style = ParagraphStyle('title', parent=styles['Heading1'],
                                      fontSize=18, textColor=colors.HexColor('#C65F2A'),
                                      spaceAfter=12, alignment=1)
        story.append(Paragraph("FEUILLE DE PRÉSENCE", title_style))
        
        # Info table
        info_data = [
            ["Classe:", session["class_name"], "Cours:", session["course_name"]],
            ["Professeur:", session["teacher_name"], "Date:", session["started_at"][:19]],
        ]
        info_table = Table(info_data, colWidths=[3*cm, 5*cm, 3*cm, 7*cm])
        info_table.setStyle(TableStyle([
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
            ('FONTNAME', (2,0), (2,-1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 10),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.5*cm))
        
        # Main attendance table
        presents = sum(1 for r in rows if r["status"] == "present")
        absents = len(rows) - presents
        taux = f"{round(presents/len(rows)*100)}%" if rows else "0%"
        
        # Stats row
        stats_data = [
            [f"✓ Presents: {presents}", f"✗ Absents: {absents}", f"Taux: {taux}"]
        ]
        stats_table = Table(stats_data, colWidths=[6*cm, 6*cm, 6*cm])
        stats_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (0,0), colors.HexColor('#E7F5ED')),
            ('BACKGROUND', (1,0), (1,0), colors.HexColor('#FDEBE8')),
            ('BACKGROUND', (2,0), (2,0), colors.HexColor('#E8EEF8')),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 11),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#E2DBD1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2DBD1')),
        ]))
        story.append(stats_table)
        story.append(Spacer(1, 0.4*cm))
        
        # Attendance table
        table_data = [["N°", "Nom de l'Étudiant", "Statut", "Heure", "Confiance"]]
        for i, row in enumerate(rows, 1):
            conf = f"{round(float(row['confidence']) * 100)}%" if row['confidence'] else "-"
            table_data.append([
                str(i), row["full_name"],
                "PRÉSENT" if row["status"] == "present" else "ABSENT",
                (row["recognized_at"] or "-")[:19],
                conf
            ])
        
        col_widths = [1.2*cm, 7*cm, 3.2*cm, 4.5*cm, 2.5*cm]
        att_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        
        style_cmds = [
            # Header
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#C65F2A')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('ALIGN', (0,0), (-1,0), 'CENTER'),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('ALIGN', (0,1), (0,-1), 'CENTER'),
            ('ALIGN', (2,1), (2,-1), 'CENTER'),
            ('ALIGN', (4,1), (4,-1), 'CENTER'),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#E2DBD1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2DBD1')),
        ]
        # Row colors
        for i, row in enumerate(rows, 1):
            bg = colors.HexColor('#E7F5ED') if row["status"] == "present" else colors.HexColor('#FDEBE8')
            style_cmds.append(('BACKGROUND', (0,i), (-1,i), bg))
        
        att_table.setStyle(TableStyle(style_cmds))
        story.append(att_table)
        
        doc.build(story)

    elif fmt == "json":
        payload = build_export_payload(session, rows)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    else:
        return jsonify({"error": "Format non supporte"}), 400

    return send_file(path, as_attachment=True, download_name=filename)


if __name__ == "__main__":

    def open_browser():
        time.sleep(2)
        webbrowser.open_new("http://127.0.0.1:5000")

    threading.Thread(target=open_browser, daemon=True).start()

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=False,
        threaded=True,
        use_reloader=False
    )
